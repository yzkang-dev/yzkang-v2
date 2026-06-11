"""
颐智康养 - 合同付款 API
"""
from datetime import date, datetime
from typing import Optional, List
from collections import defaultdict
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func, extract

from models import get_db, Contract, ContractPayment, Elder
from schemas import ContractCreate, ContractOut, ContractPaymentMark, ContractPaymentOut
from auth import RequirePermission

router = APIRouter(prefix="/api/contracts", tags=["合同付款"])


def _build_contract_out(contract: Contract) -> dict:
    """构建带统计的合同输出"""
    payments = contract.payments or []
    paid_amount = sum(p.actual_amount or 0 for p in payments)
    total_plans = len(payments)
    paid_plans = sum(1 for p in payments if p.status == "paid")
    overdue_plans = sum(1 for p in payments if p.status == "overdue")

    payments_out = [
        ContractPaymentOut.model_validate(p) for p in payments
    ]

    return {
        "id": contract.id,
        "elder_id": contract.elder_id,
        "contract_no": contract.contract_no,
        "contract_name": contract.contract_name,
        "total_amount": contract.total_amount,
        "start_date": contract.start_date,
        "end_date": contract.end_date,
        "payment_mode": contract.payment_mode,
        "status": contract.status,
        "notes": contract.notes,
        "elder_name": contract.elder.name if contract.elder else None,
        "room_number": contract.elder.room_number if contract.elder else None,
        "paid_amount": paid_amount,
        "total_plans": total_plans,
        "paid_plans": paid_plans,
        "overdue_plans": overdue_plans,
        "payments": payments_out,
    }


@router.get("/", response_model=list[ContractOut], summary="合同列表")
def list_contracts(
    elder_id: Optional[int] = Query(None, description="按老人筛选"),
    status: Optional[str] = Query(None, description="按合同状态筛选"),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    _=Depends(RequirePermission("contract:view")),
):
    """查询合同列表，支持按老人和状态筛选，自动标记逾期"""
    q = db.query(Contract).options(joinedload(Contract.elder), joinedload(Contract.payments))
    if elder_id:
        q = q.filter(Contract.elder_id == elder_id)
    if status:
        q = q.filter(Contract.status == status)
    q = q.order_by(Contract.created_at.desc())
    offset = (page - 1) * page_size
    contracts = q.offset(offset).limit(page_size).all()

    # 自动更新逾期状态
    today = date.today()
    for contract in contracts:
        if contract.status == "active":
            for p in contract.payments:
                if p.status == "pending" and p.plan_date < today:
                    p.status = "overdue"
    db.commit()

    return [_build_contract_out(c) for c in contracts]


@router.post("/", status_code=201, response_model=ContractOut, summary="创建合同")
def create_contract(
    data: ContractCreate,
    db: Session = Depends(get_db),
    _=Depends(RequirePermission("contract:create")),
):
    """创建新合同及付款计划"""
    elder = db.query(Elder).filter(Elder.id == data.elder_id).first()
    if not elder:
        raise HTTPException(status_code=404, detail="老人不存在")

    existing = db.query(Contract).filter(Contract.contract_no == data.contract_no).first()
    if existing:
        raise HTTPException(status_code=400, detail="合同编号已存在")

    contract = Contract(
        elder_id=data.elder_id,
        contract_no=data.contract_no,
        contract_name=data.contract_name,
        total_amount=data.total_amount,
        start_date=data.start_date,
        end_date=data.end_date,
        payment_mode=data.payment_mode,
        notes=data.notes,
    )
    db.add(contract)
    db.flush()

    for p in data.payments:
        payment = ContractPayment(
            contract_id=contract.id,
            payment_no=p.payment_no,
            plan_date=p.plan_date,
            plan_amount=p.plan_amount,
            notes=p.notes,
        )
        db.add(payment)

    db.commit()
    db.refresh(contract)
    # reload with relationships
    contract = db.query(Contract).options(
        joinedload(Contract.elder), joinedload(Contract.payments)
    ).filter(Contract.id == contract.id).first()
    return _build_contract_out(contract)


@router.get("/{contract_id}", response_model=ContractOut, summary="合同详情")
def get_contract(
    contract_id: int,
    db: Session = Depends(get_db),
    _=Depends(RequirePermission("contract:view")),
):
    """查看单个合同详情及付款计划"""
    contract = (
        db.query(Contract)
        .options(joinedload(Contract.elder), joinedload(Contract.payments))
        .filter(Contract.id == contract_id)
        .first()
    )
    if not contract:
        raise HTTPException(status_code=404, detail="合同不存在")
    return _build_contract_out(contract)


@router.post("/{contract_id}/payments/{payment_id}/mark", summary="标记付款")
def mark_payment(
    contract_id: int,
    payment_id: int,
    data: ContractPaymentMark,
    db: Session = Depends(get_db),
    _=Depends(RequirePermission("contract:pay")),
):
    """标记某期付款已完成，自动判断合同是否结清"""
    payment = (
        db.query(ContractPayment)
        .filter(
            ContractPayment.id == payment_id,
            ContractPayment.contract_id == contract_id,
        )
        .first()
    )
    if not payment:
        raise HTTPException(status_code=404, detail="付款记录不存在")

    if payment.status == "paid":
        raise HTTPException(status_code=400, detail="该期已付款")

    payment.actual_date = data.actual_date
    payment.actual_amount = data.actual_amount
    payment.status = "paid"
    if data.notes:
        payment.notes = data.notes

    # 检查合同所有付款是否完成
    contract = db.query(Contract).filter(Contract.id == contract_id).first()
    unpaid = (
        db.query(ContractPayment)
        .filter(
            ContractPayment.contract_id == contract_id,
            ContractPayment.status.in_(["pending", "overdue"]),
        )
        .count()
    )
    if unpaid == 0:
        contract.status = "completed"

    db.commit()
    return {"message": "付款标记成功"}


# ==================== 看板统计 API ====================

@router.get("/dashboard/stats", summary="合同统计看板")
def contract_dashboard_stats(
    db: Session = Depends(get_db),
    _=Depends(RequirePermission("contract:view")),
):
    """看板顶部统计卡片数据：合同数、金额、收款率、逾期数"""
    # 自动更新逾期
    today = date.today()
    db.query(ContractPayment).filter(
        ContractPayment.status == "pending",
        ContractPayment.plan_date < today,
    ).update({"status": "overdue"})
    db.commit()

    total_contracts = db.query(Contract).count()
    active_contracts = db.query(Contract).filter(Contract.status == "active").count()
    completed_contracts = db.query(Contract).filter(Contract.status == "completed").count()
    total_amount = db.query(func.coalesce(func.sum(Contract.total_amount), 0)).scalar()
    total_paid = db.query(func.coalesce(func.sum(ContractPayment.actual_amount), 0)).filter(
        ContractPayment.status == "paid"
    ).scalar()
    overdue_count = db.query(ContractPayment).filter(
        ContractPayment.status == "overdue"
    ).count()
    pending_count = db.query(ContractPayment).filter(
        ContractPayment.status == "pending"
    ).count()

    # 收款率
    collection_rate = round(total_paid / total_amount * 100, 1) if total_amount > 0 else 0

    return {
        "total_contracts": total_contracts,
        "active_contracts": active_contracts,
        "completed_contracts": completed_contracts,
        "total_amount": total_amount,
        "total_paid": total_paid,
        "total_unpaid": total_amount - total_paid,
        "overdue_count": overdue_count,
        "pending_count": pending_count,
        "collection_rate": collection_rate,
    }


@router.get("/dashboard/monthly-trend", summary="月度收款趋势")
def monthly_collection_trend(
    months: int = Query(12, ge=1, le=36, description="最近N个月"),
    db: Session = Depends(get_db),
    _=Depends(RequirePermission("contract:view")),
):
    """月度收款趋势图数据（最近N个月）"""
    today = date.today()
    result = []

    for i in range(months - 1, -1, -1):
        year = today.year
        month = today.month - i
        while month <= 0:
            month += 12
            year -= 1

        # 该月应收金额
        plan_total = db.query(func.coalesce(func.sum(ContractPayment.plan_amount), 0)).filter(
            extract('year', ContractPayment.plan_date) == year,
            extract('month', ContractPayment.plan_date) == month,
        ).scalar()

        # 该月实收金额
        paid_total = db.query(func.coalesce(func.sum(ContractPayment.actual_amount), 0)).filter(
            extract('year', ContractPayment.actual_date) == year,
            extract('month', ContractPayment.actual_date) == month,
            ContractPayment.status == "paid",
        ).scalar()

        result.append({
            "month": f"{year}-{month:02d}",
            "plan_amount": round(plan_total, 2),
            "paid_amount": round(paid_total, 2),
            "rate": round(paid_total / plan_total * 100, 1) if plan_total > 0 else 0,
        })

    return result


@router.get("/dashboard/overdue-analysis", summary="逾期分析")
def overdue_analysis(
    db: Session = Depends(get_db),
    _=Depends(RequirePermission("contract:view")),
):
    """逾期分析饼图数据：已付/待付/逾期分布"""
    today = date.today()
    # 自动更新逾期
    db.query(ContractPayment).filter(
        ContractPayment.status == "pending",
        ContractPayment.plan_date < today,
    ).update({"status": "overdue"})
    db.commit()

    total = db.query(ContractPayment).count()
    paid = db.query(ContractPayment).filter(ContractPayment.status == "paid").count()
    overdue = db.query(ContractPayment).filter(ContractPayment.status == "overdue").count()
    pending = db.query(ContractPayment).filter(ContractPayment.status == "pending").count()

    # 逾期金额
    overdue_amount = db.query(func.coalesce(func.sum(ContractPayment.plan_amount), 0)).filter(
        ContractPayment.status == "overdue"
    ).scalar()

    return {
        "distribution": [
            {"name": "已付", "value": paid, "color": "#52c41a"},
            {"name": "待付", "value": pending, "color": "#1890ff"},
            {"name": "逾期", "value": overdue, "color": "#ff4d4f"},
        ],
        "total": total,
        "overdue_amount": round(overdue_amount, 2),
    }


@router.get("/dashboard/kanban", summary="合同看板")
def kanban_board(
    elder_id: Optional[int] = Query(None, description="按老人筛选"),
    search: Optional[str] = Query(None, description="按合同编号/名称/老人姓名搜索"),
    db: Session = Depends(get_db),
    _=Depends(RequirePermission("contract:view")),
):
    """Kanban看板数据 - 按付款状态分列（逾期/待付/已付）"""
    today = date.today()
    db.query(ContractPayment).filter(
        ContractPayment.status == "pending",
        ContractPayment.plan_date < today,
    ).update({"status": "overdue"})
    db.commit()

    q = db.query(ContractPayment).options(
        joinedload(ContractPayment.contract).joinedload(Contract.elder)
    )

    if elder_id:
        q = q.join(Contract).filter(Contract.elder_id == elder_id)
    if search:
        q = q.join(Contract).join(Elder).filter(
            (Contract.contract_no.contains(search)) |
            (Contract.contract_name.contains(search)) |
            (Elder.name.contains(search))
        )

    payments = q.order_by(ContractPayment.plan_date).all()

    kanban = {
        "overdue": [],
        "pending": [],
        "paid": [],
    }

    for p in payments:
        contract = p.contract
        elder = contract.elder if contract else None
        item = {
            "payment_id": p.id,
            "contract_id": contract.id if contract else None,
            "contract_no": contract.contract_no if contract else "",
            "contract_name": contract.contract_name if contract else "",
            "elder_name": elder.name if elder else "",
            "room_number": elder.room_number if elder else "",
            "payment_no": p.payment_no,
            "plan_date": str(p.plan_date),
            "plan_amount": p.plan_amount,
            "actual_date": str(p.actual_date) if p.actual_date else None,
            "actual_amount": p.actual_amount,
            "status": p.status,
            "notes": p.notes,
            "total_plans": len(contract.payments) if contract else 0,
            "paid_plans": sum(1 for x in contract.payments if x.status == "paid") if contract else 0,
        }
        if p.status in kanban:
            kanban[p.status].append(item)

    return kanban


# ==================== Excel 导出 API ====================

@router.get("/export/excel", summary="导出合同Excel")
def export_contracts_excel(
    elder_id: Optional[int] = Query(None, description="按老人筛选"),
    status: Optional[str] = Query(None, description="按合同状态筛选"),
    db: Session = Depends(get_db),
    _=Depends(RequirePermission("contract:view")),
):
    """导出合同付款明细为Excel文件"""
    import io
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
    except ImportError:
        raise HTTPException(status_code=500, detail="服务端缺少openpyxl库，请先安装: pip install openpyxl")

    q = db.query(Contract).options(
        joinedload(Contract.elder), joinedload(Contract.payments)
    )
    if elder_id:
        q = q.filter(Contract.elder_id == elder_id)
    if status:
        q = q.filter(Contract.status == status)
    contracts = q.order_by(Contract.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "合同付款明细"

    # 样式
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="center", wrap_text=True)
    money_fmt = '#,##0.00'
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # 标题行
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = "颐智康养 - 合同付款明细表"
    title_cell.font = Font(name="微软雅黑", size=16, bold=True)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # 日期行
    ws.merge_cells("A2:J2")
    ws["A2"].value = f"导出时间：{datetime.now().strftime('%Y年%m月%d日 %H:%M')}"
    ws["A2"].font = Font(name="微软雅黑", size=9, color="666666")
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 20

    # 表头
    headers = [
        "合同编号", "合同名称", "老人姓名", "房间号",
        "期数", "应付日期", "应付金额",
        "实付日期", "实付金额", "状态",
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[4].height = 28

    # 列宽
    col_widths = [16, 22, 12, 10, 8, 14, 14, 14, 14, 10]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # 数据行
    row = 5
    status_fill_map = {
        "paid": PatternFill(start_color="F6FFED", end_color="F6FFED", fill_type="solid"),
        "overdue": PatternFill(start_color="FFF2F0", end_color="FFF2F0", fill_type="solid"),
        "pending": PatternFill(start_color="E6F7FF", end_color="E6F7FF", fill_type="solid"),
    }
    status_label_map = {"paid": "已付", "overdue": "逾期", "pending": "待付"}

    for contract in contracts:
        elder = contract.elder
        for p in (contract.payments or []):
            ws.cell(row=row, column=1, value=contract.contract_no).alignment = cell_align
            ws.cell(row=row, column=2, value=contract.contract_name).alignment = cell_align
            ws.cell(row=row, column=3, value=elder.name if elder else "").alignment = cell_align
            ws.cell(row=row, column=4, value=elder.room_number if elder else "").alignment = cell_align
            ws.cell(row=row, column=5, value=f"第{p.payment_no}期").alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=row, column=6, value=str(p.plan_date)).alignment = cell_align
            ws.cell(row=row, column=7, value=p.plan_amount).number_format = money_fmt
            ws.cell(row=row, column=7).alignment = Alignment(horizontal="right", vertical="center")
            ws.cell(row=row, column=8, value=str(p.actual_date) if p.actual_date else "").alignment = cell_align
            ws.cell(row=row, column=9, value=p.actual_amount if p.actual_amount else 0).number_format = money_fmt
            ws.cell(row=row, column=9).alignment = Alignment(horizontal="right", vertical="center")

            status_cell = ws.cell(row=row, column=10, value=status_label_map.get(p.status, p.status))
            status_cell.alignment = Alignment(horizontal="center", vertical="center")
            if p.status in status_fill_map:
                status_cell.fill = status_fill_map[p.status]

            for c in range(1, 11):
                ws.cell(row=row, column=c).border = thin_border
            row += 1

    # 合计行
    if contracts:
        total_plan = sum(p.plan_amount for c in contracts for p in (c.payments or []))
        total_actual = sum(p.actual_amount or 0 for c in contracts for p in (c.payments or []) if p.status == "paid")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        summary_cell = ws.cell(row=row, column=1, value="合计")
        summary_cell.font = Font(name="微软雅黑", bold=True, size=11)
        summary_cell.alignment = Alignment(horizontal="center", vertical="center")

        plan_cell = ws.cell(row=row, column=7, value=round(total_plan, 2))
        plan_cell.number_format = money_fmt
        plan_cell.font = Font(bold=True)
        plan_cell.alignment = Alignment(horizontal="right", vertical="center")

        actual_cell = ws.cell(row=row, column=9, value=round(total_actual, 2))
        actual_cell.number_format = money_fmt
        actual_cell.font = Font(bold=True)
        actual_cell.alignment = Alignment(horizontal="right", vertical="center")

        for c in range(1, 11):
            ws.cell(row=row, column=c).border = thin_border
            ws.cell(row=row, column=c).fill = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")

    # 冻结窗格
    ws.freeze_panes = "A5"

    # 输出到内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    from urllib.parse import quote
    filename = f"颐智康养_合同付款明细_{datetime.now().strftime('%Y%m%d')}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )
